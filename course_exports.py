"""CourseKraft PDF and editable PowerPoint export helpers."""

from html import escape as html_escape
from io import BytesIO
import textwrap


def _chunks(values, size):
    return [values[index:index + size] for index in range(0, len(values), size)]


def _short_text(value, limit):
    value = str(value or "")
    return value if len(value) <= limit else value[:limit - 1].rstrip() + "…"


def create_pptx(course):
    """Build a designed deck while keeping long-form material in speaker notes."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    ink = RGBColor(32, 33, 59)
    deep_ink = RGBColor(25, 26, 48)
    violet = RGBColor(101, 84, 217)
    coral = RGBColor(229, 178, 160)
    ivory = RGBColor(245, 240, 232)
    white = RGBColor(255, 255, 255)
    muted = RGBColor(104, 104, 124)
    pale = RGBColor(251, 249, 245)

    def textbox(slide, x, y, width, height, text, size=18, color=ink, bold=False, font="Aptos"):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = Inches(0.06)
        frame.margin_right = Inches(0.06)
        frame.margin_top = Inches(0.03)
        frame.margin_bottom = Inches(0.02)
        paragraph = frame.paragraphs[0]
        paragraph.text = str(text)
        paragraph.font.name = font
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        return shape

    def rectangle(slide, x, y, width, height, fill, rounded=False):
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.fill.background()
        return shape

    def add_notes(slide, text):
        if text:
            slide.notes_slide.notes_text_frame.text = str(text)

    def base_slide(title, section, dark=False):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = deep_ink if dark else pale
        if not dark:
            rectangle(slide, 0, 0, 13.333, 0.13, violet)
            textbox(slide, 0.65, 0.35, 11.8, 0.3, section.upper(), 9, violet, True)
            textbox(slide, 0.65, 0.75, 12, 0.65, title, 27, ink, True)
            rectangle(slide, 0.65, 1.48, 1.0, 0.055, coral)
            textbox(slide, 0.65, 7.08, 9, 0.2, str(course.get("title", "CourseKraft AI"))[:85], 8, muted)
            textbox(slide, 12.1, 7.04, 0.55, 0.25, f"{len(presentation.slides):02d}", 9, muted, True)
        return slide

    # Title page
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = deep_ink
    rectangle(slide, 0, 0, 0.18, 7.5, coral)
    rectangle(slide, 9.25, 0, 4.1, 7.5, ink)
    rectangle(slide, 10.0, 0.85, 2.1, 0.12, violet)
    rectangle(slide, 10.6, 1.18, 1.5, 0.12, coral)
    textbox(slide, 0.85, 0.75, 7.8, 0.35, "COURSEKRAFT AI  /  TEACHING DECK", 11, coral, True)
    textbox(slide, 0.85, 1.55, 8.0, 2.1, course.get("title", "Course Learning Deck"), 34, white, True)
    metadata = "  ·  ".join(filter(None, [
        str(course.get("level", "")),
        f"{course['durationMinutes']} minutes" if course.get("durationMinutes") else "",
        str(course.get("language", "")),
    ]))
    textbox(slide, 0.9, 4.05, 7.3, 0.8, metadata, 16, RGBColor(225, 221, 238))
    textbox(slide, 0.9, 5.45, 7.5, 0.6, "Concepts  ·  Examples  ·  Practice", 13, coral, True)
    textbox(slide, 10.0, 5.8, 2.8, 0.6, "LEARN  /  APPLY  /  BUILD", 11, white, True)

    objectives = course.get("objectives", [])
    for page, group in enumerate(_chunks(objectives, 6), start=1):
        slide = base_slide("What you will learn", "Learning outcomes")
        for index, objective in enumerate(group):
            y = 1.85 + index * 0.78
            ordinal = (page - 1) * 6 + index + 1
            rectangle(slide, 0.75, y, 0.55, 0.52, violet, True)
            textbox(slide, 0.75, y + 0.08, 0.55, 0.28, f"{ordinal:02d}", 12, white, True)
            rectangle(slide, 1.5, y, 11.05, 0.52, white, True)
            textbox(slide, 1.78, y + 0.08, 10.5, 0.38, _short_text(objective, 110), 12, ink)
        add_notes(slide, "LEARNING OBJECTIVES\n" + "\n".join(objectives))

    for index, item in enumerate(course.get("slides", []), start=1):
        bullets = item.get("bullets", [])
        bullet_pages = _chunks(bullets, 6) or [[]]
        for page_index, bullet_group in enumerate(bullet_pages, start=1):
            title = item.get("title", f"Concept {index}")
            if page_index > 1:
                title = f"{title} · continued"
            slide = base_slide(title, f"Module {index:02d}")
            columns = 2 if len(bullet_group) > 3 else 1
            rows = (len(bullet_group) + columns - 1) // columns
            card_width = 5.85 if columns == 2 else 11.9
            for bullet_index, bullet in enumerate(bullet_group):
                column = bullet_index // rows if columns == 2 else 0
                row = bullet_index % rows if columns == 2 else bullet_index
                x = 0.7 + column * 6.0
                y = 1.85 + row * (4.75 / max(rows, 1))
                height = min(1.08, 4.65 / max(rows, 1) - 0.1)
                rectangle(slide, x, y, card_width, height, white, True)
                rectangle(slide, x, y, 0.09, height, coral if bullet_index % 2 else violet)
                font_size = 16 if len(str(bullet)) < 115 else 13
                textbox(slide, x + 0.22, y + 0.12, card_width - 0.42, height - 0.18,
                        "•  " + _short_text(bullet, 200), font_size, ink)
            visual_labels = item.get("visualLabels", [])
            note_lines = ["SLIDE CONTENT", "\n".join(bullets), "TRAINER NOTES", item.get("speakerNotes", "")]
            if item.get("visualType"):
                note_lines.extend([f"Visual type: {item['visualType']}",
                                   f"Visual description: {item.get('visualDescription', '')}",
                                   "Visual labels: " + ", ".join(visual_labels)])
            add_notes(slide, "\n\n".join(line for line in note_lines if line))

        visual_type = str(item.get("visualType", "none")).lower()
        labels = item.get("visualLabels", [])
        if visual_type not in ("none", "image") and labels:
            diagram = base_slide(f"Visual guidance · {item.get('title', f'Concept {index}')}", "Concept visual")
            description = item.get("visualDescription", "")
            if description:
                textbox(diagram, 0.75, 1.65, 11.8, 0.55, description, 13, muted)
            horizontal = visual_type in ("flowchart", "timeline") and len(labels) <= 6
            if horizontal:
                count = len(labels)
                card_width = min(2.25, 10.9 / max(count, 1))
                gap = (11.5 - count * card_width) / max(count - 1, 1)
                for label_index, label in enumerate(labels):
                    x = 0.9 + label_index * (card_width + gap)
                    rectangle(diagram, x, 3.05, card_width, 1.15, white, True)
                    textbox(diagram, x + 0.12, 3.21, card_width - 0.24, 0.82, _short_text(label, 50), 12, ink, True)
                    if label_index < count - 1:
                        textbox(diagram, x + card_width, 3.4, gap, 0.35, "→", 22, violet, True)
            else:
                for label_index, label in enumerate(labels):
                    column, row = label_index % 2, label_index // 2
                    x, y = 0.9 + column * 5.95, 2.45 + row * 0.94
                    rectangle(diagram, x, y, 5.55, 0.72, white, True)
                    rectangle(diagram, x, y, 0.08, 0.72, coral if label_index % 2 else violet)
                    textbox(diagram, x + 0.22, y + 0.13, 5.1, 0.45, _short_text(label, 70), 12, ink, True)
            add_notes(diagram, "Visual description: " + description + "\n\nSuggested visual labels:\n" + "\n".join(labels))

    for example in course.get("codeExamples", []):
        code = example.get("code", "")
        lines = code.splitlines() or ["No code supplied for this worked example."]
        visual_lines = []
        for line in lines:
            visual_lines.extend(textwrap.wrap(
                line, width=125, replace_whitespace=False, drop_whitespace=False,
                break_long_words=True, break_on_hyphens=False,
            ) or [""])
        lines = visual_lines
        pages = _chunks(lines, 20) or [[]]
        for page_index, page_lines in enumerate(pages, start=1):
            slide = base_slide(example.get("title", "Worked example"), "Code example", dark=True)
            suffix = "" if page_index == 1 else f" · continued {page_index}"
            textbox(slide, 0.7, 0.45, 11.8, 0.3,
                    str(example.get("language", course.get("language", "CODE"))).upper() + "  /  WALKTHROUGH",
                    10, coral, True)
            textbox(slide, 0.7, 0.9, 11.8, 0.6, str(example.get("title", "Worked example")) + suffix, 23, white, True)
            rectangle(slide, 0.7, 1.75, 11.9, 4.45, ink, True)
            textbox(slide, 0.95, 1.98, 11.35, 4.05, "\n".join(page_lines), 10, RGBColor(238, 234, 244), False, "Consolas")
            if page_index == len(pages):
                textbox(slide, 0.8, 6.4, 11.8, 0.55,
                        "Expected output: " + _short_text(example.get("expectedOutput", "Not specified"), 180),
                        11, coral, True)
            add_notes(slide, "\n\n".join(filter(None, [
                "COMPLETE SOURCE CODE\n" + code,
                f"Test description: {example.get('testDescription', '')}",
                f"Expected output: {example.get('expectedOutput', '')}",
            ])))

    exercises = course.get("exercises", [])
    for group_index, group in enumerate(_chunks(exercises, 4), start=1):
        slide = base_slide("Practice exercises", "Apply")
        notes = ["FULL EXERCISE DETAILS"]
        for index, exercise in enumerate(group):
            column, row = index % 2, index // 2
            x, y = 0.75 + column * 6.0, 1.85 + row * 2.35
            rectangle(slide, x, y, 5.65, 2.05, white, True)
            textbox(slide, x + 0.2, y + 0.12, 5.2, 0.3,
                    exercise.get("title", f"Exercise {index + 1}"), 15, violet, True)
            textbox(slide, x + 0.2, y + 0.5, 5.2, 1.0,
                    _short_text(exercise.get("description", ""), 300), 11, ink)
            footer = " · ".join(filter(None, [exercise.get("difficulty", ""),
                                                 _short_text(exercise.get("testHint", ""), 80)]))
            textbox(slide, x + 0.2, y + 1.58, 5.2, 0.3, footer, 9, muted)
            notes.append(f"{exercise.get('title', 'Exercise')}\nDifficulty: {exercise.get('difficulty', '')}\n"
                         f"{exercise.get('description', '')}\nAcceptance hint: {exercise.get('testHint', '')}")
        add_notes(slide, "\n\n".join(notes))

    quiz = course.get("quiz", [])
    for question_index, question in enumerate(quiz, start=1):
        slide = base_slide("Knowledge check", "Review")
        textbox(slide, 0.85, 1.9, 11.6, 0.8,
                f"{question_index}. {_short_text(question.get('question', ''), 190)}", 16, ink, True)
        for option_index, option in enumerate(question.get("options", [])):
            column, row = option_index % 2, option_index // 2
            x, y = 0.95 + column * 5.9, 3.0 + row * 0.7
            rectangle(slide, x, y, 5.55, 0.52, white, True)
            textbox(slide, x + 0.15, y + 0.09, 5.2, 0.34,
                    f"{chr(65 + option_index)}. {_short_text(option, 75)}", 10, ink)
        options = question.get("options", [])
        correct_index = question.get("correctIndex")
        answer = options[correct_index] if isinstance(correct_index, int) and 0 <= correct_index < len(options) else "Not supplied"
        full_options = "\n".join(f"{chr(65 + index)}. {option}" for index, option in enumerate(options))
        add_notes(slide, f"INSTRUCTOR ANSWER KEY\nQuestion: {question.get('question', '')}\n{full_options}\n"
                         f"Answer: {answer}\nExplanation: {question.get('explanation', '')}")

    terms = course.get("cheatSheet", [])
    valid_terms = [t for t in terms if isinstance(t, dict)]
    for group in _chunks(valid_terms, 10):
        slide = base_slide("Quick reference", "Revision")
        notes = ["FULL QUICK REFERENCE"]
        for index, term in enumerate(group):
            column, row = index % 2, index // 2
            x, y = 0.7 + column * 6.0, 1.82 + row * 0.98
            rectangle(slide, x, y, 5.8, 0.82, white, True)
            textbox(slide, x + 0.2, y + 0.1, 5.35, 0.25, _short_text(term.get("term", "Term"), 50), 11, violet, True)
            textbox(slide, x + 0.2, y + 0.39, 5.35, 0.32, _short_text(term.get("definition", ""), 80), 9, ink)
            notes.append(f"{term.get('term', '')}: {term.get('definition', '')}")
        add_notes(slide, "\n\n".join(notes))

    slide = base_slide("Apply it: your next step", "Wrap-up", dark=True)
    textbox(slide, 0.9, 1.3, 11.5, 0.8, "From understanding to doing", 30, white, True)
    textbox(slide, 0.95, 2.5, 10.8, 1.2,
            "Choose one concept from this session. Explain it in your own words, then apply it in a small example or exercise.",
            21, RGBColor(235, 231, 243))
    rectangle(slide, 0.95, 4.45, 4.1, 0.08, coral)
    textbox(slide, 0.95, 4.85, 10.8, 0.6, "Questions?  ·  Discussion  ·  Practice", 16, coral, True)

    output = BytesIO()
    presentation.save(output)
    output.seek(0)
    return output


def create_pdf(course):
    """Build a paginated, escaped teaching pack as a PDF."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        Preformatted,
        SimpleDocTemplate,
        Spacer,
    )

    def safe_text(value):
        return html_escape(str(value), quote=False)

    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="CourseKraftTitle", parent=styles["Title"],
        textColor=colors.HexColor("#20213b"), spaceAfter=14,
    ))
    styles.add(ParagraphStyle(
        name="CourseKraftCode", parent=styles["Code"], fontName="Courier",
        fontSize=8, leading=10, splitLongWords=1, wordWrap="CJK",
        backColor=colors.HexColor("#f5f0e8"), borderPadding=6,
    ))
    styles.add(ParagraphStyle(
        name="CourseKraftMetadata", parent=styles["Normal"],
        textColor=colors.HexColor("#56566a"), spaceAfter=8,
    ))

    metadata = "  ·  ".join(filter(None, [
        f"Topic: {course.get('topic', '')}" if course.get("topic") else "",
        f"Level: {course.get('level', '')}" if course.get("level") else "",
        f"Duration: {course.get('durationMinutes')} minutes" if course.get("durationMinutes") else "",
        f"Language: {course.get('language', '')}" if course.get("language") else "",
    ]))
    story = [
        Paragraph(safe_text(course.get("title", "CourseKraft AI")), styles["CourseKraftTitle"]),
        Paragraph("COURSEKRAFT AI  /  TEACHING PACK", styles["CourseKraftMetadata"]),
    ]
    if metadata:
        story.append(Paragraph(safe_text(metadata), styles["CourseKraftMetadata"]))

    objectives = course.get("objectives", [])
    if objectives:
        story.append(Paragraph("Learning objectives", styles["Heading1"]))
        for objective in objectives:
            story.append(Paragraph("• " + safe_text(objective), styles["BodyText"]))

    slides = course.get("slides", [])
    if slides:
        story.append(Paragraph("Lesson slides and trainer notes", styles["Heading1"]))
    for index, slide in enumerate(slides, start=1):
        story.append(Paragraph(f"{index}. {safe_text(slide.get('title', 'Slide'))}", styles["Heading2"]))
        for bullet in slide.get("bullets", []):
            story.append(Paragraph("• " + safe_text(bullet), styles["BodyText"]))
        if slide.get("speakerNotes"):
            story.append(Paragraph("<b>Trainer notes:</b> " + safe_text(slide["speakerNotes"]), styles["BodyText"]))
        if slide.get("visualType") or slide.get("visualDescription") or slide.get("visualLabels"):
            story.append(Paragraph("<b>Visual guidance:</b> " + safe_text(slide.get("visualType", "suggestion")), styles["BodyText"]))
            if slide.get("visualDescription"):
                story.append(Paragraph(safe_text(slide["visualDescription"]), styles["BodyText"]))
            if slide.get("visualLabels"):
                story.append(Paragraph("Labels: " + safe_text(" · ".join(slide["visualLabels"])), styles["BodyText"]))
        story.append(Spacer(1, 8))

    examples = course.get("codeExamples", [])
    if examples:
        story.append(PageBreak())
        story.append(Paragraph("Worked examples", styles["Heading1"]))
    for example in examples:
        story.append(Paragraph(safe_text(example.get("title", "Worked example")), styles["Heading2"]))
        if example.get("language"):
            story.append(Paragraph(safe_text(example["language"].upper()), styles["CourseKraftMetadata"]))
        if example.get("code"):
            story.append(Preformatted(safe_text(example["code"]), styles["CourseKraftCode"], maxLineLength=100))
        if example.get("testDescription"):
            story.append(Paragraph("<b>Test description:</b> " + safe_text(example["testDescription"]), styles["BodyText"]))
        if example.get("expectedOutput"):
            story.append(Paragraph("<b>Expected output (not verified):</b> " + safe_text(example["expectedOutput"]), styles["BodyText"]))

    exercises = course.get("exercises", [])
    if exercises:
        story.append(PageBreak())
        story.append(Paragraph("Exercises", styles["Heading1"]))
    for exercise in exercises:
        story.append(Paragraph(safe_text(exercise.get("title", "Exercise")), styles["Heading2"]))
        if exercise.get("difficulty"):
            story.append(Paragraph("Difficulty: " + safe_text(exercise["difficulty"]), styles["CourseKraftMetadata"]))
        if exercise.get("description"):
            story.append(Paragraph(safe_text(exercise["description"]), styles["BodyText"]))
        if exercise.get("testHint"):
            story.append(Paragraph("<b>Acceptance hint:</b> " + safe_text(exercise["testHint"]), styles["BodyText"]))

    questions = course.get("quiz", [])
    if questions:
        story.append(PageBreak())
        story.append(Paragraph("Quiz and instructor answer key", styles["Heading1"]))
    for index, question in enumerate(questions, start=1):
        story.append(Paragraph(f"{index}. {safe_text(question.get('question', 'Question'))}", styles["Heading2"]))
        for option_index, option in enumerate(question.get("options", [])):
            story.append(Paragraph(f"{chr(65 + option_index)}. {safe_text(option)}", styles["BodyText"]))
        options = question.get("options", [])
        correct_index = question.get("correctIndex")
        if isinstance(correct_index, int) and 0 <= correct_index < len(options):
            story.append(Paragraph("<b>Instructor answer:</b> " + safe_text(options[correct_index]), styles["BodyText"]))
        if question.get("explanation"):
            story.append(Paragraph("<b>Explanation:</b> " + safe_text(question["explanation"]), styles["BodyText"]))

    terms = course.get("cheatSheet", [])
    if terms:
        story.append(PageBreak())
        story.append(Paragraph("Quick reference", styles["Heading1"]))
    for term in terms:
        if not isinstance(term, dict):
            continue
        story.append(Paragraph(safe_text(term.get("term", "Term")), styles["Heading2"]))
        if term.get("category"):
            story.append(Paragraph("Category: " + safe_text(term["category"]), styles["CourseKraftMetadata"]))
        story.append(Paragraph(safe_text(term.get("definition", "")), styles["BodyText"]))
        if term.get("example"):
            story.append(Preformatted(safe_text(term["example"]), styles["CourseKraftCode"], maxLineLength=100))
        if term.get("related"):
            related_str = " · ".join(str(r) for r in term["related"] if r)
            if related_str:
                story.append(Paragraph("<b>Connected concepts:</b> " + safe_text(related_str), styles["CourseKraftMetadata"]))
        story.append(Spacer(1, 6))

    document.build(story)
    output.seek(0)
    return output
