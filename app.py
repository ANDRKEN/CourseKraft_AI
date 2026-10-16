import os, json, re, urllib.request, urllib.error, tempfile, time
from flask import Flask, render_template, request, jsonify, send_file
from dotenv import load_dotenv
load_dotenv()
from io import BytesIO
from datetime import datetime

app = Flask(__name__)
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")
JUDGE0_URL = os.getenv("JUDGE0_URL", "").strip().rstrip("/")
JUDGE0_API_KEY = os.getenv("JUDGE0_API_KEY", "").strip()

SAMPLE_COURSE = {
    "title": "Object-Oriented Programming in Java",
    "topic": "Object-Oriented Programming in Java",
    "level": "Beginner",
    "durationMinutes": 45,
    "language": "Java",
    "objectives": [
        "Explain classes and objects",
        "Describe constructors and encapsulation",
        "Apply object-oriented concepts in a small Java program"
    ],
    "slides": [
        {"title":"What is OOP?","bullets":["A programming approach built around objects","Objects combine state (data) and behaviour (methods)","Java supports classes, inheritance, polymorphism and encapsulation"],"speakerNotes":"Start with a real-world example such as a Student object."},
        {"title":"Classes and Objects","bullets":["A class is a blueprint","An object is an instance of a class","Fields represent state; methods represent behaviour"],"speakerNotes":"Compare a class to a blueprint and objects to actual houses built from it."},
        {"title":"Encapsulation","bullets":["Keep data and behaviour together","Use access modifiers to control access","Expose safe operations through methods"],"speakerNotes":"Explain why fields are often private and accessed through methods."},
        {"title":"Constructors","bullets":["A constructor initializes a new object","Its name matches the class name","It has no return type"],"speakerNotes":"Show a constructor that initializes a student's name and roll number."},
        {"title":"Key Takeaways","bullets":["Classes define structure","Objects are runtime instances","Encapsulation protects internal state","Practise by modelling a real-world entity"],"speakerNotes":"Ask students to describe one class they would create for a library system."}
    ],
    "codeExamples": [
        {"id":"java-001","title":"Class and object example","language":"java","code":"public class Main {\n  static class Student {\n    String name;\n    Student(String name) { this.name = name; }\n  }\n  public static void main(String[] args) {\n    Student s = new Student(\"Asha\");\n    System.out.println(s.name);\n  }\n}","testDescription":"Expected output: Asha","expectedOutput":"Asha"}
    ],
    "exercises": [
        {"title":"Create a Book class","description":"Create a Book class with title and author fields, then create one object and print its title.","difficulty":"Easy","testHint":"The output should include the book title."},
        {"title":"Add a constructor","description":"Add a constructor that initializes the title and author fields.","difficulty":"Medium","testHint":"Create the object using the constructor and print both fields."}
    ],
    "quiz": [
        {"question":"What is a class in Java?","options":["A blueprint for objects","A running process","A database table","A loop"],"correctIndex":0,"explanation":"A class defines the fields and methods that its objects can have."},
        {"question":"Which statement best describes an object?","options":["A package of source files","An instance of a class","A compiler error","A primitive type"],"correctIndex":1,"explanation":"An object is an instance created from a class."},
        {"question":"What is the main purpose of encapsulation?","options":["Make every field public","Control access to internal state","Remove methods","Replace constructors"],"correctIndex":1,"explanation":"Encapsulation groups state and behaviour and controls access to implementation details."}
    ],
    "cheatSheet": [
        {"term":"Class","definition":"A blueprint defining fields and methods."},
        {"term":"Object","definition":"An instance of a class."},
        {"term":"Constructor","definition":"A special member used to initialize a new object."},
        {"term":"Encapsulation","definition":"Bundling data and methods while controlling access."},
        {"term":"Method","definition":"A named block of code that performs an operation."}
    ]
}


def call_gemini(prompt):
    if not GEMINI_API_KEY:
        return None, "No GEMINI_API_KEY configured."

    body = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {
            "temperature": 0.4,
            "responseMimeType": "application/json",
            "maxOutputTokens": 6000
        }
    }

    # A 503 can persist for one model while another supported Gemini model is
    # healthy. Try the configured model first, then stable fallbacks. Users can
    # override the fallback order with GEMINI_FALLBACK_MODELS in .env.
    configured_fallbacks = os.getenv("GEMINI_FALLBACK_MODELS", "")
    fallbacks = [m.strip() for m in configured_fallbacks.split(",") if m.strip()]
    if not fallbacks:
        fallbacks = ["gemini-3.5-flash-lite", "gemini-3.8-flash"]
    models = list(dict.fromkeys([GEMINI_MODEL] + fallbacks))
    retryable_statuses = {408, 429, 500, 502, 503, 504}
    failover_statuses = {500, 502, 503, 504}
    last_error = "Gemini API could not generate a response."

    for model_index, model in enumerate(models):
        max_attempts = 3 if model_index == 0 else 2
        url = ("https://generativelanguage.googleapis.com/v1beta/"
               f"models/{model}:generateContent?key={GEMINI_API_KEY}")
        for attempt in range(max_attempts):
            req = urllib.request.Request(
                url,
                data=json.dumps(body).encode("utf-8"),
                headers={"Content-Type": "application/json"},
                method="POST"
            )
            try:
                with urllib.request.urlopen(req, timeout=45) as resp:
                    data = json.loads(resp.read().decode("utf-8"))

                text = data["candidates"][0]["content"]["parts"][0]["text"].strip()
                if text.startswith("```"):
                    text = text.split("\n", 1)[1].rsplit("```", 1)[0].strip()
                parsed = json.loads(text)

                required = ["title", "slides", "codeExamples", "exercises", "quiz", "cheatSheet"]
                if not all(key in parsed for key in required):
                    raise ValueError("AI response is missing required course fields.")

                parsed.setdefault("topic", parsed["title"])
                parsed.setdefault("level", "Beginner")
                parsed.setdefault("durationMinutes", 45)
                parsed.setdefault("language", "Java")
                parsed.setdefault("objectives", [])
                if model_index:
                    return parsed, f"Generated with Gemini fallback model {model}; {GEMINI_MODEL} was temporarily unavailable."
                return parsed, f"Generated with Gemini API ({model})."

            except urllib.error.HTTPError as e:
                details = e.read().decode("utf-8", errors="replace")
                last_error = f"Gemini API HTTP {e.code} on {model}: {details[:500]}"
                if e.code == 429:
                    try:
                        error_data = json.loads(details).get("error", {})
                    except (ValueError, AttributeError):
                        error_data = {}
                    provider_message = str(error_data.get("message", ""))
                    quota_error = (
                        "quota" in provider_message.lower()
                        or "generate_content_free_tier" in details.lower()
                        or "quota_exceeded" in str(error_data.get("status", "")).lower()
                    )
                    if quota_error:
                        # Quotas are model-specific. Try the configured fallback
                        # models once, without sleeping or retrying this quota.
                        last_error = f"Gemini quota exhausted on {model}."
                        if model_index < len(models) - 1:
                            break
                        limit = re.search(r"limit:\s*(\d+)", provider_message, re.IGNORECASE)
                        retry = re.search(r"retry in\s+([0-9a-z.]+)", provider_message, re.IGNORECASE)
                        limit_text = f" (limit {limit.group(1)})" if limit else ""
                        retry_text = f" Try again in about {retry.group(1).rstrip('.')}." if retry else " Wait for the quota reset shown in Google AI Studio."
                        return None, (f"Gemini's quota for {model}{limit_text} is exhausted.{retry_text} "
                                      "See your project's usage and limits at https://ai.dev/rate-limit.")
                # A 503 explicitly means this model is overloaded. Repeating
                # the same request only adds delay, so switch models at once.
                if e.code == 503 and model_index < len(models) - 1:
                    break
                if e.code in retryable_statuses and attempt < max_attempts - 1:
                    retry_after = e.headers.get("Retry-After", "")
                    try:
                        delay = min(8.0, max(0.5, float(retry_after))) if retry_after else 2 ** attempt
                    except ValueError:
                        delay = 2 ** attempt
                    time.sleep(delay)
                    continue
                if (e.code in failover_statuses or (model_index > 0 and e.code == 404)) and model_index < len(models) - 1:
                    break
                return None, f"{last_error} Tried {model_index + 1} model(s); please retry shortly."
            except (TimeoutError, urllib.error.URLError) as e:
                if attempt < max_attempts - 1:
                    time.sleep(2 ** attempt)
                    continue
                return None, f"Gemini connection error after {max_attempts} attempts: {str(e)[:400]}"
            except Exception as e:
                return None, f"Gemini API error: {str(e)[:400]}"

    return None, f"{last_error} All configured Gemini models are temporarily unavailable."



@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/status")
def status():
    return jsonify({
        "geminiConfigured": bool(GEMINI_API_KEY),
        "judge0Configured": bool(JUDGE0_URL),
        "model": GEMINI_MODEL,
        "message": "API keys are read from local environment variables."
    })


@app.route("/api/sample")
def sample():
    """Return the built-in demo course without calling Gemini."""
    course = {**SAMPLE_COURSE, "_source": "Sample course"}
    return jsonify({"course": course, "message": "Loaded the built-in sample course.",
                    "source": "Sample course"})


@app.route("/api/generate", methods=["POST"])
def generate():
    data = request.get_json(silent=True) or {}
    topic = str(data.get("topic", "")).strip()[:180]
    level = str(data.get("level", "Beginner")).strip()[:50]
    language = str(data.get("language", "Java")).strip()[:40]
    try:
        duration = max(10, min(240, int(data.get("durationMinutes", 45))))
    except (TypeError, ValueError):
        duration = 45

    if not topic:
        return jsonify({"error": "Enter a topic before generating a course."}), 400
    if not GEMINI_API_KEY:
        return jsonify({
            "error": "Gemini API key is not configured. Add GEMINI_API_KEY to your .env file and restart the app."
        }), 503

    prompt = f"""
You are CourseKraft AI, an expert instructional designer and subject-matter tutor.
Create a genuinely topic-specific teaching package. Do not reuse generic sample material.
Topic: {topic}
Audience/level: {level}
Programming language (when relevant): {language}
Session duration: {duration} minutes.

Return ONLY valid JSON with exactly these top-level keys:
"title", "topic", "level", "durationMinutes", "language", "objectives",
"slides", "codeExamples", "exercises", "quiz", "cheatSheet".

Requirements:
- Create 7 to 8 useful slides, ordered as a coherent lesson from introduction to recap.
- Each slide must have "title", "bullets" (2 to 4 informative, specific bullets), and
  "speakerNotes" (a useful trainer script of about 40-70 words).
- Include practical examples, definitions, step-by-step explanations, and real-world applications.
- For every slide also include "visualType" (one of "image", "flowchart", "architecture",
  "timeline", "comparison", "chart", "none"), "visualDescription" (specific instructions
  for a relevant visual), and "visualLabels" (array of labels for diagrams; otherwise []).
- Include 1 to 2 "codeExamples" if programming is relevant. Each has "id", "title",
  "language", "code", "testDescription", "expectedOutput". For non-programming topics,
  provide 1-2 worked examples instead with the same fields and empty code if not applicable.
- Include 2 exercises, each with "title", "description", "difficulty", "testHint".
- Include 3 quiz questions, each with "question", four "options", "correctIndex" (0-3),
  and "explanation".
- Include 5 to 6 cheatSheet items with "term" and "definition".
- Be accurate, avoid fabricated citations/statistics, and flag uncertainty instead of inventing facts.
- Keep every field concise and specific. Do not add extra sections or repeat content. Do not put markdown fences around the JSON.
"""
    course, message = call_gemini(prompt)
    if not course:
        return jsonify({"error": message, "source": "Gemini AI", "sampleFallback": False}), 502

    course["_source"] = "Gemini AI"
    return jsonify({"course": course, "message": message, "source": "Gemini AI"})


@app.route("/api/verify", methods=["POST"])
def verify():
    data = request.get_json(force=True)
    language = str(data.get("language","java")).lower()
    code = str(data.get("code",""))[:12000]
    expected = str(data.get("expectedOutput","")).strip()
    if not code.strip():
        return jsonify({"status":"NOT VERIFIED","message":"No code was supplied."})
    if not JUDGE0_URL:
        return jsonify({
            "status":"NOT VERIFIED",
            "message":"No isolated code runner is configured. The code was not executed.",
            "details":"Configure JUDGE0_URL (and JUDGE0_API_KEY if required) to enable real execution. This is intentionally not reported as a pass."
        })
    lang_ids = {"python":71, "python3":71, "java":62, "javascript":63, "js":63, "cpp":54, "c++":54}
    lang_id = lang_ids.get(language)
    if not lang_id:
        return jsonify({"status":"NOT VERIFIED","message":f"Unsupported language: {language}"})
    payload = {"source_code":code, "language_id":lang_id, "stdin":""}
    headers = {"Content-Type":"application/json"}
    if JUDGE0_API_KEY:
        headers["X-RapidAPI-Key"] = JUDGE0_API_KEY
        # Some hosted deployments use a host header; allow user to set the full service URL.
    try:
        req = urllib.request.Request(JUDGE0_URL + "/submissions?base64_encoded=false&wait=true",
                                     data=json.dumps(payload).encode(), headers=headers, method="POST")
        with urllib.request.urlopen(req, timeout=35) as resp:
            result = json.loads(resp.read().decode())
        status_desc = (result.get("status") or {}).get("description","")
        stdout = (result.get("stdout") or "").strip()
        stderr = (result.get("stderr") or result.get("compile_output") or "").strip()
        if status_desc == "Accepted" and (not expected or stdout == expected):
            state = "PASSED"
            msg = "Execution completed successfully" + (" and output matched." if expected else ".")
        elif status_desc == "Accepted":
            state = "FAILED"
            msg = "Code executed, but output did not match the expected output."
        elif status_desc in ("Wrong Answer", "Compilation Error", "Runtime Error", "Time Limit Exceeded"):
            state = "FAILED"
            msg = status_desc
        else:
            state = "NOT VERIFIED"
            msg = "Runner returned an inconclusive result: " + (status_desc or "unknown status")
        return jsonify({"status":state,"message":msg,"runnerStatus":status_desc,
                        "stdout":stdout,"stderr":stderr,"expectedOutput":expected,
                        "executionTime":result.get("time"),"memory":result.get("memory")})
    except Exception as e:
        return jsonify({"status":"NOT VERIFIED","message":"Execution service could not verify the code.",
                        "details":str(e)[:240]})

@app.route("/api/export/pptx", methods=["POST"])
def export_pptx():
    course = request.get_json(force=True)
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.enum.shapes import MSO_SHAPE

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        NAVY = RGBColor(17, 24, 49)
        INK = RGBColor(31, 41, 68)
        PURPLE = RGBColor(105, 85, 230)
        TEAL = RGBColor(39, 190, 180)
        WHITE = RGBColor(255, 255, 255)
        MUTED = RGBColor(105, 116, 140)
        PALE = RGBColor(244, 246, 252)

        def textbox(slide, x, y, w, h, text, size=18, color=INK, bold=False, font="Aptos"):
            shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            tf = shape.text_frame
            tf.clear(); tf.word_wrap = True
            tf.margin_left = Inches(0.04); tf.margin_right = Inches(0.04)
            tf.margin_top = Inches(0.03); tf.margin_bottom = Inches(0.02)
            p = tf.paragraphs[0]; p.text = str(text)
            p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
            return shape

        def rect(slide, x, y, w, h, fill, radius=False):
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                           Inches(x), Inches(y), Inches(w), Inches(h))
            shape.fill.solid(); shape.fill.fore_color.rgb = fill
            shape.line.fill.background()
            return shape

        def base_slide(title, section, n, dark=False):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = NAVY if dark else PALE
            if not dark:
                rect(slide, 0, 0, 13.333, 0.13, PURPLE)
                textbox(slide, .65, .35, 11.8, .3, section.upper(), 9, PURPLE, True)
                textbox(slide, .65, .75, 12, .65, title, 27, NAVY, True)
                rect(slide, .65, 1.48, 1.0, .055, TEAL)
                textbox(slide, .65, 7.08, 9, .2, str(course.get("title", "CourseKraft AI"))[:85], 8, MUTED)
                textbox(slide, 12.1, 7.04, .55, .25, f"{n:02d}", 9, MUTED, True)
            return slide

        # Designed title slide
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
        rect(slide, 0, 0, .18, 7.5, TEAL)
        rect(slide, 9.25, 0, 4.1, 7.5, INK)
        rect(slide, 10.0, .85, 2.1, .12, PURPLE)
        rect(slide, 10.6, 1.18, 1.5, .12, TEAL)
        textbox(slide, .85, .75, 7.8, .35, "COURSEKRAFT AI  /  TRAINING DECK", 11, TEAL, True)
        textbox(slide, .85, 1.55, 8.0, 2.1, course.get("title", "Course Learning Deck"), 34, WHITE, True)
        textbox(slide, .9, 4.05, 7.3, .8, f"{course.get('level','Beginner')}  •  {course.get('durationMinutes',45)} minutes  •  {course.get('language','')}", 16, RGBColor(205, 213, 235))
        textbox(slide, .9, 5.45, 7.5, .6, "Concepts  ·  Worked examples  ·  Practice", 13, TEAL, True)
        textbox(slide, 10.0, 5.8, 2.8, .6, "LEARN  /  APPLY  /  BUILD", 11, WHITE, True)

        # Learning objectives slide
        objectives = course.get("objectives", [])
        if objectives:
            slide = base_slide("What you will learn", "Learning outcomes", 2)
            for i, obj in enumerate(objectives[:6]):
                y = 1.85 + i * .78
                rect(slide, .75, y, .55, .52, PURPLE, True)
                textbox(slide, .75, y+.08, .55, .28, f"{i+1:02d}", 12, WHITE, True)
                rect(slide, 1.5, y, 11.05, .52, WHITE, True)
                textbox(slide, 1.78, y+.08, 10.5, .38, obj, 16, INK, False)

        slides = course.get("slides", [])
        for idx, item in enumerate(slides, 1):
            slide = base_slide(item.get("title", f"Concept {idx}"), f"Module {idx:02d}", len(prs.slides)+1)
            bullets = item.get("bullets", [])
            # Balanced two-column cards for dense content
            cols = 2 if len(bullets) > 3 else 1
            rows = (len(bullets) + cols - 1) // cols
            card_w = 5.85 if cols == 2 else 11.9
            for j, bullet in enumerate(bullets[:6]):
                col = j // rows if cols == 2 else 0
                row = j % rows if cols == 2 else j
                x = .7 + col * 6.0
                y = 1.85 + row * (4.75 / max(rows, 1))
                h = min(1.08, 4.65 / max(rows, 1) - .1)
                rect(slide, x, y, card_w, h, WHITE, True)
                rect(slide, x, y, .09, h, TEAL if j % 2 else PURPLE)
                textbox(slide, x+.22, y+.12, card_w-.42, h-.18, "•  " + str(bullet), 16 if len(str(bullet)) < 115 else 13, INK)
            notes = str(item.get("speakerNotes", "")).strip()
            if notes:
                try:
                    slide.notes_slide.notes_text_frame.text = "TRAINER NOTES\n\n" + notes
                except Exception:
                    pass

            # Turn the model's diagram labels into native editable slide shapes.
            # This produces genuine PPT visuals without relying on remote images.
            visual_type = str(item.get("visualType", "none")).lower()
            labels = [str(label)[:90] for label in item.get("visualLabels", []) if str(label).strip()][:8]
            if visual_type not in ("none", "image") and labels:
                diagram = base_slide(f"{item.get('title', 'Concept')} — visual", "Concept visual", len(prs.slides)+1)
                description = str(item.get("visualDescription", ""))[:220]
                if description:
                    textbox(diagram, .75, 1.65, 11.8, .55, description, 13, MUTED)
                horizontal = visual_type in ("flowchart", "timeline") and len(labels) <= 6
                if horizontal:
                    count = len(labels)
                    card_w = min(2.25, 10.9 / max(count, 1))
                    gap = (11.5 - count * card_w) / max(count - 1, 1)
                    y = 3.05
                    for j, label in enumerate(labels):
                        x = .9 + j * (card_w + gap)
                        rect(diagram, x, y, card_w, 1.15, WHITE, True)
                        textbox(diagram, x+.12, y+.16, card_w-.24, .82, label, 13, INK, True)
                        if j < count - 1:
                            textbox(diagram, x+card_w, y+.35, gap, .35, "→", 22, PURPLE, True)
                else:
                    cols = 2
                    for j, label in enumerate(labels):
                        col, row = j % cols, j // cols
                        x, y = .9 + col * 5.95, 2.45 + row * .94
                        rect(diagram, x, y, 5.55, .72, WHITE, True)
                        rect(diagram, x, y, .08, .72, TEAL if j % 2 else PURPLE)
                        textbox(diagram, x+.22, y+.13, 5.1, .45, label, 14, INK, True)

        # Code examples as readable dedicated slides
        for code in course.get("codeExamples", []):
            slide = base_slide(code.get("title", "Worked code example"), "Worked example", len(prs.slides)+1, dark=True)
            textbox(slide, .7, .45, 11.8, .3, str(code.get("language", course.get("language", "CODE"))).upper()+"  /  WALKTHROUGH", 10, TEAL, True)
            textbox(slide, .7, .9, 11.8, .6, code.get("title", "Worked example"), 25, WHITE, True)
            rect(slide, .7, 1.75, 11.9, 4.55, INK, True)
            code_text = str(code.get("code", ""))
            textbox(slide, .95, 1.98, 11.35, 4.05, code_text[:2600], 11, RGBColor(224, 231, 247), False, "Consolas")
            textbox(slide, .8, 6.45, 11.8, .5, "Expected output: " + str(code.get("expectedOutput", "Not specified"))[:220], 12, TEAL, True)
            try:
                slide.notes_slide.notes_text_frame.text = "Test description: " + str(code.get("testDescription", "Not specified"))
            except Exception:
                pass

        # Practice materials are part of the teaching package and belong in the deck.
        exercises = course.get("exercises", [])
        if exercises:
            slide = base_slide("Practice exercises", "Apply", len(prs.slides)+1)
            for i, exercise in enumerate(exercises[:6]):
                col, row = i % 2, i // 2
                x, y = .75 + col * 6.0, 1.85 + row * 1.65
                rect(slide, x, y, 5.65, 1.42, WHITE, True)
                textbox(slide, x+.2, y+.12, 5.2, .3, exercise.get("title", f"Exercise {i+1}"), 15, PURPLE, True)
                textbox(slide, x+.2, y+.48, 5.2, .7, exercise.get("description", ""), 11, INK)

        quiz = course.get("quiz", [])
        for offset in range(0, len(quiz), 2):
            slide = base_slide("Knowledge check", "Review", len(prs.slides)+1)
            for j, question in enumerate(quiz[offset:offset+2]):
                y = 1.85 + j * 2.35
                textbox(slide, .85, y, 11.6, .58, f"{offset+j+1}. {question.get('question', '')}", 17, NAVY, True)
                options = question.get("options", [])
                for k, option in enumerate(options[:4]):
                    col, row = k % 2, k // 2
                    x, oy = .95 + col * 5.9, y + .72 + row * .6
                    rect(slide, x, oy, 5.55, .46, WHITE, True)
                    textbox(slide, x+.15, oy+.08, 5.2, .3, f"{chr(65+k)}. {option}", 11, INK)
            try:
                answers = [f"{offset+i+1}. {q.get('options', [])[q.get('correctIndex', -1)]}" for i, q in enumerate(quiz[offset:offset+2]) if isinstance(q.get("correctIndex"), int) and 0 <= q.get("correctIndex") < len(q.get("options", []))]
                if answers:
                    slide.notes_slide.notes_text_frame.text = "INSTRUCTOR ANSWER KEY\n" + "\n".join(answers)
            except Exception:
                pass

        # Recap / key terms slide
        terms = course.get("cheatSheet", [])
        if terms:
            slide = base_slide("Quick reference", "Revision", len(prs.slides)+1)
            for i, term in enumerate(terms[:10]):
                col = i % 2; row = i // 2
                x = .7 + col * 6.0; y = 1.82 + row * .98
                rect(slide, x, y, 5.8, .82, WHITE, True)
                textbox(slide, x+.2, y+.1, 5.35, .25, term.get("term", "Term"), 13, PURPLE, True)
                textbox(slide, x+.2, y+.39, 5.35, .32, term.get("definition", ""), 10, INK)

        # Closing slide
        slide = base_slide("Apply it: your next step", "Wrap-up", len(prs.slides)+1, dark=True)
        textbox(slide, .9, 1.3, 11.5, .8, "From understanding to doing", 30, WHITE, True)
        textbox(slide, .95, 2.5, 10.8, 1.2, "Choose one concept from this session. Explain it in your own words, then apply it in a small example or exercise.", 21, RGBColor(220, 226, 243))
        rect(slide, .95, 4.45, 4.1, .08, TEAL)
        textbox(slide, .95, 4.85, 10.8, .6, "Questions?  •  Discussion  •  Practice", 16, TEAL, True)
        out = BytesIO()
        prs.save(out); out.seek(0)
        filename = "coursekraft_designed_course.pptx"
        return send_file(out, as_attachment=True, download_name=filename,
                         mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation")
    except Exception as e:
        return jsonify({"error":"PPTX export failed. Install requirements and retry: "+str(e)[:200]}), 500

@app.route("/api/export/pdf", methods=["POST"])
def export_pdf():
    course = request.get_json(force=True)
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, ListFlowable, ListItem
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        from reportlab.lib.units import mm
        out = BytesIO()
        doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(name="CFTitle", parent=styles["Title"], textColor=colors.HexColor("#5b4bdb"), spaceAfter=14))
        story = [Paragraph(str(course.get("title","CourseKraft AI")), styles["CFTitle"]),
                 Paragraph(f"Level: {course.get('level','')} | Duration: {course.get('durationMinutes',45)} minutes", styles["Normal"]),
                 Spacer(1,12)]
        for s in course.get("slides",[]):
            story.append(Paragraph(str(s.get("title","Slide")), styles["Heading2"]))
            for b in s.get("bullets",[]):
                story.append(Paragraph("• " + str(b), styles["BodyText"]))
            if s.get("speakerNotes"):
                story.append(Paragraph("<b>Trainer notes:</b> " + str(s["speakerNotes"]), styles["BodyText"]))
            story.append(Spacer(1,8))
        story.append(PageBreak())
        story.append(Paragraph("Code Examples", styles["Heading1"]))
        for c in course.get("codeExamples",[]):
            story.append(Paragraph(str(c.get("title","Code example")), styles["Heading2"]))
            code_text = str(c.get("code","")).replace("&","&amp;").replace("<","&lt;").replace(">","&gt;").replace("\n","<br/>")
            story.append(Paragraph(code_text, styles["Code"]))
            story.append(Paragraph("Expected output: " + str(c.get("expectedOutput","Not specified")), styles["BodyText"]))
        story.append(PageBreak())
        story.append(Paragraph("Exercises", styles["Heading1"]))
        for ex in course.get("exercises",[]):
            story.append(Paragraph(str(ex.get("title","Exercise")), styles["Heading2"]))
            story.append(Paragraph(str(ex.get("description","")), styles["BodyText"]))
        story.append(Paragraph("Quiz", styles["Heading1"]))
        for i,q in enumerate(course.get("quiz",[]),1):
            story.append(Paragraph(f"{i}. {q.get('question','')}", styles["BodyText"]))
            for j,opt in enumerate(q.get("options",[])):
                story.append(Paragraph(f"   {chr(65+j)}. {opt}", styles["BodyText"]))
        story.append(Paragraph("Cheat Sheet", styles["Heading1"]))
        for item in course.get("cheatSheet",[]):
            story.append(Paragraph(f"<b>{item.get('term','')}</b>: {item.get('definition','')}", styles["BodyText"]))
        doc.build(story); out.seek(0)
        return send_file(out, as_attachment=True, download_name="coursekraft_course.pdf", mimetype="application/pdf")
    except Exception as e:
        return jsonify({"error":"PDF export failed. Install requirements and retry: "+str(e)[:200]}), 500

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.getenv("PORT","5000")), debug=False)
