"""Regression tests for the existing Flask application."""

import io
import json
from email.message import Message
import unittest
from copy import deepcopy
from unittest import mock
from urllib.error import HTTPError

import dotenv

# Importing the app normally loads .env. Patch that call so tests never read it.
with mock.patch.object(dotenv, "load_dotenv"):
    import app as course_app


SAMPLE_COURSE = {
    "title": "Intro to Python",
    "topic": "Python basics",
    "level": "Beginner",
    "durationMinutes": 30,
    "language": "Python",
    "objectives": ["Print a value"],
    "slides": [{
        "title": "Hello",
        "bullets": ["Use print() to display text."],
        "speakerNotes": "Ask learners what the program displays.",
        "visualType": "flowchart",
        "visualDescription": "Show the input moving through print to the output.",
        "visualLabels": ["Input", "print()", "Output"],
    }],
    "codeExamples": [{
        "title": "Hello world",
        "language": "python",
        "code": "print('Hello, world!')",
        "testDescription": "Print a greeting",
        "expectedOutput": "Hello, world!",
    }],
    "exercises": [{
        "title": "Change the greeting",
        "description": "Print a different greeting.",
        "difficulty": "Easy",
        "testHint": "Check the displayed text.",
    }],
    "quiz": [{
        "question": "Which function displays text?",
        "options": ["print", "input", "range", "len"],
        "correctIndex": 0,
        "explanation": "print() writes output.",
    }],
    "cheatSheet": [{"term": "print", "definition": "Displays a value."}],
}


def gemini_response(course):
    response = mock.MagicMock()
    response.__enter__.return_value.read.return_value = json.dumps({
        "candidates": [{"content": {"parts": [{"text": json.dumps(course)}]}}]
    }).encode("utf-8")
    return response


def runner_response(result):
    response = mock.MagicMock()
    response.__enter__.return_value.read.return_value = json.dumps(result).encode("utf-8")
    return response


class FlaskAppTests(unittest.TestCase):
    def setUp(self):
        course_app.app.config.update(TESTING=True)
        self.client = course_app.app.test_client()

    def test_homepage_returns_success(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"<html", response.data.lower())

    def test_status_returns_valid_json(self):
        response = self.client.get("/api/status")
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertIsInstance(data["geminiConfigured"], bool)
        self.assertIsInstance(data["judge0Configured"], bool)
        self.assertIsInstance(data["model"], str)

    def test_generation_without_provider_is_unavailable_without_network(self):
        with mock.patch.object(course_app, "GEMINI_API_KEY", ""), \
             mock.patch.object(course_app.gemini_provider.urllib.request, "urlopen") as urlopen:
            response = self.client.post("/api/generate", json={
                "topic": "Testing fallback",
                "level": "Intermediate",
                "durationMinutes": 45,
                "language": "Python",
            })

        self.assertEqual(response.status_code, 503)
        result = response.get_json()
        self.assertEqual(result["source"], "unavailable")
        self.assertEqual(result["generationStatus"], "unavailable")
        self.assertIsNone(result["course"])
        self.assertIn("not configured", result["message"].lower())
        urlopen.assert_not_called()

    def test_sample_course_is_explicit_and_keeps_its_own_topic(self):
        response = self.client.get("/api/sample")

        self.assertEqual(response.status_code, 200)
        result = response.get_json()
        self.assertEqual(result["source"], "Sample content")
        self.assertEqual(result["generationStatus"], "sample")
        self.assertEqual(result["course"]["topic"], "Object-Oriented Programming in Java")
        self.assertEqual(result["course"]["_source"], "Sample content")

    def test_successful_provider_response_is_marked_as_ai_generated(self):
        provider_course = dict(SAMPLE_COURSE)
        with mock.patch.object(course_app, "GEMINI_API_KEY", "configured-for-test"), \
             mock.patch.object(course_app, "call_gemini", return_value=(provider_course, "Generated with Gemini API.")):
            response = self.client.post("/api/generate", json={"topic": "Python basics"})

        self.assertEqual(response.status_code, 200)
        result = response.get_json()
        self.assertEqual(result["source"], "Gemini AI")
        self.assertEqual(result["generationStatus"], "generated")
        self.assertEqual(result["course"]["_source"], "Gemini AI")

    def test_mocked_generation_preserves_rich_schema_and_request_settings(self):
        provider_course = deepcopy(SAMPLE_COURSE)
        provider_course.update({
            "topic": "wrong provider topic",
            "level": "wrong provider level",
            "durationMinutes": 30,
            "language": "Java",
        })
        provider_response = gemini_response(provider_course)
        with mock.patch.object(course_app, "GEMINI_API_KEY", "configured-for-test"), \
             mock.patch.object(course_app, "GEMINI_FALLBACK_MODELS", ("gemini-3.8-flash",)), \
             mock.patch.object(course_app.gemini_provider.urllib.request, "urlopen", return_value=provider_response) as urlopen:
            response = self.client.post("/api/generate", json={
                "topic": "Binary trees",
                "level": "Intermediate",
                "durationMinutes": 60,
                "language": "Python",
            })

        self.assertEqual(response.status_code, 200)
        course = response.get_json()["course"]
        self.assertEqual(course["topic"], "Binary trees")
        self.assertEqual(course["level"], "Intermediate")
        self.assertEqual(course["durationMinutes"], 60)
        self.assertEqual(course["language"], "Python")
        self.assertEqual(course["slides"][0]["visualLabels"], ["Input", "print()", "Output"])
        self.assertEqual(course["codeExamples"][0]["testDescription"], "Print a greeting")
        self.assertEqual(course["quiz"][0]["explanation"], "print() writes output.")
        self.assertEqual(course["cheatSheet"][0]["definition"], "Displays a value.")
        request = urlopen.call_args.args[0]
        self.assertNotIn("configured-for-test", request.full_url)
        self.assertEqual(request.get_header("X-goog-api-key"), "configured-for-test")

    def test_generation_prompt_includes_course_settings_and_visual_schema(self):
        prompt = course_app.build_generation_prompt("Photosynthesis", "Beginner", 35, "Java")

        for phrase in ("Topic: Photosynthesis", "Learner level: Beginner", "35 minutes",
                       "Preferred programming language: Java", "visualDescription", "visualLabels",
                       "expectedOutput", "do not claim the example was executed"):
            with self.subTest(phrase=phrase):
                self.assertIn(phrase.lower(), prompt.lower())

    def test_gemini_capacity_failure_falls_back_without_leaking_key(self):
        provider_response = gemini_response(SAMPLE_COURSE)
        unavailable = HTTPError(
            "https://generativelanguage.googleapis.com/", 503,
            "upstream unavailable", Message(), None,
        )
        with mock.patch.object(course_app, "GEMINI_API_KEY", "fake-provider-key"), \
             mock.patch.object(course_app, "GEMINI_FALLBACK_MODELS", ("gemini-3.8-flash",)), \
             mock.patch.object(course_app.gemini_provider.urllib.request, "urlopen", side_effect=[unavailable, provider_response]) as urlopen, \
             mock.patch.object(course_app.gemini_provider.time, "sleep"):
            course, message = course_app.call_gemini("safe prompt")

        self.assertIsNotNone(course)
        self.assertIn("Gemini", message)
        self.assertEqual(urlopen.call_count, 2)
        self.assertIn("gemini-3.8-flash", urlopen.call_args_list[1].args[0].full_url)
        self.assertNotIn("fake-provider-key", " ".join(call.args[0].full_url for call in urlopen.call_args_list))

    def test_gemini_nonretryable_provider_error_is_safe(self):
        denied = HTTPError(
            "https://generativelanguage.googleapis.com/", 403,
            "fake-provider-key should not be shown", Message(), None,
        )
        with mock.patch.object(course_app, "GEMINI_API_KEY", "fake-provider-key"), \
             mock.patch.object(course_app, "GEMINI_FALLBACK_MODELS", ()), \
             mock.patch.object(course_app.gemini_provider.urllib.request, "urlopen", side_effect=denied) as urlopen:
            course, message = course_app.call_gemini("safe prompt")

        self.assertIsNone(course)
        self.assertIn("rejected", message.lower())
        self.assertNotIn("fake-provider-key", message)
        urlopen.assert_called_once()

    def test_gemini_retries_a_transient_provider_failure_once(self):
        transient = HTTPError(
            "https://generativelanguage.googleapis.com/", 500,
            "temporary provider issue", Message(), None,
        )
        provider_response = gemini_response(SAMPLE_COURSE)
        with mock.patch.object(course_app, "GEMINI_API_KEY", "configured-for-test"), \
             mock.patch.object(course_app, "GEMINI_FALLBACK_MODELS", ()), \
             mock.patch.object(course_app.gemini_provider.urllib.request, "urlopen", side_effect=[transient, provider_response]) as urlopen, \
             mock.patch.object(course_app.gemini_provider.time, "sleep") as sleep:
            course, _ = course_app.call_gemini("safe prompt")

        self.assertIsNotNone(course)
        self.assertEqual(urlopen.call_count, 2)
        self.assertEqual(urlopen.call_args_list[0].args[0].full_url, urlopen.call_args_list[1].args[0].full_url)
        sleep.assert_called_once()

    def test_long_topic_is_truncated_correctly(self):
        provider_response = gemini_response(SAMPLE_COURSE)
        with mock.patch.object(course_app, "GEMINI_API_KEY", "configured-for-test"), \
             mock.patch.object(course_app, "build_generation_prompt", wraps=course_app.build_generation_prompt) as mock_prompt, \
             mock.patch.object(course_app.gemini_provider.urllib.request, "urlopen", return_value=provider_response):

            long_topic = "x" * 200
            response = self.client.post("/api/generate", json={"topic": long_topic})

            self.assertEqual(response.status_code, 200)
            mock_prompt.assert_called_once()
            called_topic = mock_prompt.call_args[0][0]
            self.assertEqual(len(called_topic), 180)
            self.assertEqual(called_topic, "x" * 180)

    def test_invalid_generation_inputs_return_client_errors(self):
        invalid_requests = [
            ("not-json", True),
            ([], False),
            ({"topic": 42}, False),
            ({"topic": "valid", "durationMinutes": "many"}, False),
            ({"topic": "valid", "level": ["Beginner"]}, False),
            ({"topic": "valid", "language": 123}, False),
            ({"topic": "valid", "durationMinutes": True}, False),
        ]
        for payload, malformed_json in invalid_requests:
            with self.subTest(payload=payload):
                if malformed_json:
                    response = self.client.post(
                        "/api/generate", data=payload, content_type="application/json"
                    )
                else:
                    response = self.client.post("/api/generate", json=payload)
                self.assertEqual(response.status_code, 400)
                self.assertIn("error", response.get_json())

    def test_oversized_request_gets_json_client_error(self):
        response = self.client.post(
            "/api/export/pdf",
            data=b"{" + b" " * (2 * 1024 * 1024) + b"}",
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 413)
        self.assertIn("2 MB", response.get_json()["error"])

    def test_generation_rejects_malformed_provider_course(self):
        malformed_course = deepcopy(SAMPLE_COURSE)
        malformed_course["slides"][0]["bullets"] = "not a list"
        with mock.patch.object(
            course_app, "call_gemini", return_value=(malformed_course, "Generated with Gemini API.")
        ):
            response = self.client.post("/api/generate", json={"topic": "Python basics"})

        self.assertEqual(response.status_code, 502)
        result = response.get_json()
        self.assertIsNone(result["course"])
        self.assertEqual(result["generationStatus"], "invalid")
        self.assertIn("bullets", result["error"])

    def test_gemini_response_is_validated_without_network_access(self):
        malformed_course = deepcopy(SAMPLE_COURSE)
        malformed_course["quiz"][0]["correctIndex"] = "zero"
        provider_response = mock.MagicMock()
        provider_response.__enter__.return_value.read.return_value = json.dumps({
            "candidates": [{"content": {"parts": [{"text": json.dumps(malformed_course)}]}}]
        }).encode("utf-8")

        with mock.patch.object(course_app, "GEMINI_API_KEY", "test-key-not-a-real-credential"), \
             mock.patch.object(course_app.gemini_provider.urllib.request, "urlopen", return_value=provider_response) as urlopen:
            course, message = course_app.call_gemini("safe test prompt")

        self.assertIsNone(course)
        self.assertIn("correctIndex", message)
        urlopen.assert_called_once()

    def test_malformed_gemini_json_is_reported_as_invalid_content(self):
        provider_response = mock.MagicMock()
        provider_response.__enter__.return_value.read.return_value = json.dumps({
            "candidates": [{"content": {"parts": [{"text": "{not valid JSON"}]}}]
        }).encode("utf-8")

        with mock.patch.object(course_app, "GEMINI_API_KEY", "test-key-not-a-real-credential"), \
             mock.patch.object(course_app.gemini_provider.urllib.request, "urlopen", return_value=provider_response) as urlopen:
            course, message = course_app.call_gemini("safe test prompt")

        self.assertIsNone(course)
        self.assertIn("invalid course content", message)
        urlopen.assert_called_once()

    def test_verification_without_runner_never_claims_pass(self):
        with mock.patch.object(course_app, "JUDGE0_URL", ""), \
             mock.patch.object(course_app.judge0_provider.urllib.request, "urlopen") as urlopen:
            response = self.client.post("/api/verify", json={
                "language": "python",
                "code": "print('safe test payload')",
                "expectedOutput": "safe test payload",
            })

        self.assertEqual(response.status_code, 200)
        result = response.get_json()
        self.assertEqual(result["status"], "NOT VERIFIED")
        self.assertNotEqual(result["status"], "PASSED")
        urlopen.assert_not_called()

    def test_pdf_export_returns_nonempty_pdf(self):
        response = self.client.post("/api/export/pdf", json=SAMPLE_COURSE)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/pdf")
        self.assertTrue(response.data.startswith(b"%PDF-"))
        self.assertGreater(len(response.data), 500)

    def test_powerpoint_export_returns_valid_presentation(self):
        response = self.client.post("/api/export/pptx", json=SAMPLE_COURSE)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
        self.assertGreater(len(response.data), 500)
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(response.data))
        self.assertGreater(len(presentation.slides), 0)
        slide_text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
        notes_text = "\n".join(slide.notes_slide.notes_text_frame.text for slide in presentation.slides)
        self.assertIn("COURSEKRAFT AI", slide_text)
        self.assertIn("Input", slide_text)
        self.assertIn("Test description: Print a greeting", notes_text)
        self.assertIn("Ask learners what the program displays.", notes_text)
        self.assertIn("Check the displayed text.", notes_text)
        self.assertIn("INSTRUCTOR ANSWER KEY", notes_text)
        self.assertIn("Displays a value.", notes_text)

    def test_runner_success_requires_accepted_execution_and_matching_output(self):
        response = runner_response({
            "status": {"description": "Accepted"},
            "stdout": "Hello, world!\n",
            "stderr": "",
            "time": "0.01",
            "memory": 1234,
        })
        with mock.patch.object(course_app, "JUDGE0_URL", "https://judge0-ce.p.rapidapi.com"), \
             mock.patch.object(course_app, "JUDGE0_API_KEY", "runner-key-for-test"), \
             mock.patch.object(course_app, "JUDGE0_API_KEY_HEADER", ""), \
             mock.patch.object(course_app, "JUDGE0_LANGUAGE_IDS", ""), \
             mock.patch.object(course_app.judge0_provider.urllib.request, "urlopen", return_value=response) as urlopen:
            result = self.client.post("/api/verify", json={
                "language": "python", "code": "print('Hello, world!')", "expectedOutput": "Hello, world!",
            }).get_json()

        self.assertEqual(result["status"], "PASSED")
        self.assertIn("judge0-ce.p.rapidapi.com", result["runnerDescription"])
        self.assertEqual(result["runnerStatus"], "Accepted")
        self.assertEqual(result["stdout"], "Hello, world!\n")
        request = urlopen.call_args.args[0]
        self.assertNotIn("runner-key-for-test", request.full_url)
        self.assertEqual(request.get_header("X-rapidapi-key"), "runner-key-for-test")
        self.assertEqual(request.get_header("X-rapidapi-host"), "judge0-ce.p.rapidapi.com")

    def test_runner_mismatch_and_execution_failures_are_not_passes(self):
        cases = [
            ({"status": {"description": "Accepted"}, "stdout": "different"}, "FAILED"),
            ({"status": {"description": "Compilation Error"}, "compile_output": "compiler detail"}, "FAILED"),
            ({"status": {"description": "In Queue"}}, "NOT VERIFIED"),
        ]
        for payload, expected_status in cases:
            with self.subTest(expected_status=expected_status, runner_status=payload["status"]["description"]):
                with mock.patch.object(course_app, "JUDGE0_URL", "https://runner.example"), \
                     mock.patch.object(course_app, "JUDGE0_LANGUAGE_IDS", ""), \
                     mock.patch.object(course_app.judge0_provider.urllib.request, "urlopen", return_value=runner_response(payload)):
                    response = self.client.post("/api/verify", json={
                        "language": "python", "code": "print('safe')", "expectedOutput": "expected",
                    })
                self.assertEqual(response.get_json()["status"], expected_status)

    def test_runner_failure_is_safe_and_not_verified(self):
        with mock.patch.object(course_app, "JUDGE0_URL", "https://runner.example"), \
             mock.patch.object(course_app, "JUDGE0_LANGUAGE_IDS", ""), \
             mock.patch.object(course_app.judge0_provider.urllib.request, "urlopen", side_effect=RuntimeError("fake-token-must-not-leak")):
            result = self.client.post("/api/verify", json={
                "language": "python", "code": "print('safe')",
            }).get_json()

        self.assertEqual(result["status"], "NOT VERIFIED")
        self.assertNotIn("fake-token-must-not-leak", json.dumps(result))

    def test_runner_supports_only_configured_language_ids(self):
        with mock.patch.object(course_app, "JUDGE0_URL", "https://runner.example"), \
             mock.patch.object(course_app, "JUDGE0_LANGUAGE_IDS", '{"python":71}'), \
             mock.patch.object(course_app.judge0_provider.urllib.request, "urlopen") as urlopen:
            result = self.client.post("/api/verify", json={
                "language": "java", "code": "class Main {}",
            }).get_json()

        self.assertEqual(result["status"], "NOT VERIFIED")
        urlopen.assert_not_called()

    def test_runner_uses_a_configured_provider_specific_auth_header(self):
        with mock.patch.object(course_app, "JUDGE0_URL", "https://runner.example"), \
             mock.patch.object(course_app, "JUDGE0_API_KEY", "runner-token-for-test"), \
             mock.patch.object(course_app, "JUDGE0_API_KEY_HEADER", "Authorization"), \
             mock.patch.object(course_app, "JUDGE0_LANGUAGE_IDS", '{"python":71}'), \
             mock.patch.object(course_app.judge0_provider.urllib.request, "urlopen", return_value=runner_response({
                 "status": {"description": "Accepted"}, "stdout": "ok",
             })) as urlopen:
            result = self.client.post("/api/verify", json={"language": "python", "code": "print('ok')"}).get_json()

        self.assertEqual(result["status"], "PASSED")
        request = urlopen.call_args.args[0]
        self.assertEqual(request.get_header("Authorization"), "runner-token-for-test")
        self.assertNotIn("runner-token-for-test", request.full_url)

    def test_verification_rejects_invalid_and_overlong_payloads(self):
        payloads = [
            [],
            {"language": "python", "code": []},
            {"language": "python", "code": "x" * (course_app.MAX_CODE_LENGTH + 1)},
            {"language": "python", "code": "x", "expectedOutput": []},
            {"language": "python", "code": "x", "expectedOutput": "x" * (course_app.MAX_EXPECTED_OUTPUT_LENGTH + 1)},
        ]
        for payload in payloads:
            with self.subTest(payload_type=type(payload).__name__):
                response = self.client.post("/api/verify", json=payload)
                self.assertEqual(response.status_code, 400)
                self.assertIn("error", response.get_json())

    def test_export_routes_reject_invalid_payloads(self):
        for path in ("/api/export/pdf", "/api/export/pptx"):
            for payload in ([], {"slides": "not-a-list"}, {"slides": ["not-an-object"]}):
                with self.subTest(path=path, payload=payload):
                    response = self.client.post(path, json=payload)
                    self.assertEqual(response.status_code, 400)
                    self.assertIn("error", response.get_json())

    def test_export_routes_reject_invalid_nested_types_and_limits(self):
        invalid_courses = [
            {**SAMPLE_COURSE, "slides": [{"title": "Slide", "bullets": [3]}]},
            {**SAMPLE_COURSE, "quiz": [{"question": "Q", "options": ["A", "B"], "correctIndex": True}]},
            {**SAMPLE_COURSE, "title": None},
            {**SAMPLE_COURSE, "title": "x" * 181},
            {**SAMPLE_COURSE, "slides": [{"title": "Slide", "bullets": []}] * 61},
        ]
        for path in ("/api/export/pdf", "/api/export/pptx"):
            for course in invalid_courses:
                with self.subTest(path=path, course=course):
                    response = self.client.post(path, json=course)
                    self.assertEqual(response.status_code, 400)
                    self.assertIn("error", response.get_json())

    def test_export_routes_keep_optional_course_fields_optional(self):
        for path in ("/api/export/pdf", "/api/export/pptx"):
            with self.subTest(path=path):
                response = self.client.post(path, json={"slides": []})
                self.assertEqual(response.status_code, 200)

    def test_rich_visual_metadata_is_preserved_and_validated(self):
        course = deepcopy(SAMPLE_COURSE)
        self.assertIsNone(course_app.validate_course(course))
        self.assertEqual(course["slides"][0]["visualType"], "flowchart")
        self.assertEqual(course["slides"][0]["visualDescription"], "Show the input moving through print to the output.")
        self.assertEqual(course["slides"][0]["visualLabels"], ["Input", "print()", "Output"])

        course["slides"][0]["visualLabels"] = "not-a-list"
        self.assertIn("visualLabels", course_app.validate_course(course))

    def test_pdf_escapes_markup_in_course_content(self):
        course = deepcopy(SAMPLE_COURSE)
        course["title"] = "Intro to <b>Python</b> & safety"
        course["slides"][0]["speakerNotes"] = "Treat <script>alert(1)</script> as text."

        response = self.client.post("/api/export/pdf", json=course)

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data.startswith(b"%PDF-"))

    def test_creation_page_has_one_provider_notice_and_an_explicit_sample_action(self):
        response = self.client.get("/")
        page = response.get_data(as_text=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(page.count('id="providerNotice"'), 1)
        self.assertNotIn('id="creationAlert"', page)
        self.assertIn('id="loadSampleBtn"', page)
        self.assertIn("No AI call is made when you open the sample.", page)


if __name__ == "__main__":
    unittest.main()
